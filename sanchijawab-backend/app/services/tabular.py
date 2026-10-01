"""CSV/TSV/XLSX ingestion (FR-K6 follow-up).

Spreadsheets need table-aware chunking, not the word-count window used for
prose (chunker.py): slicing a markdown table by word count splits rows
across chunk boundaries and loses the header/value association entirely.
Instead, each row becomes its own self-contained "field: value" chunk, so a
single row is always retrieved and read as one complete record.
"""
from __future__ import annotations

import csv
import io

from openpyxl import load_workbook

from .parser import extension_of

CSV_EXTENSIONS = {".csv", ".tsv"}
XLSX_EXTENSIONS = {".xlsx"}
TABULAR_EXTENSIONS = CSV_EXTENSIONS | XLSX_EXTENSIONS

MAX_ROWS = 5000  # sane cap so a bad upload can't hang ingestion indefinitely


def _row_to_chunk(sheet_label: str, row_num: int, headers: list[str], row: list) -> str | None:
    fields = [
        f"{h.strip()}: {str(v).strip()}"
        for h, v in zip(headers, row)
        if h and h.strip() and v is not None and str(v).strip()
    ]
    if not fields:
        return None
    return f"{sheet_label}, row {row_num}:\n" + "\n".join(fields)


def _parse_csv_bytes(filename: str, data: bytes, delimiter: str) -> list[str]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = list(reader)
    if not rows:
        return []

    headers = rows[0]
    chunks = []
    for i, row in enumerate(rows[1 : MAX_ROWS + 1], start=2):
        chunk = _row_to_chunk(filename, i, headers, row)
        if chunk:
            chunks.append(chunk)
    return chunks


def _parse_xlsx_bytes(filename: str, data: bytes) -> list[str]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    chunks: list[str] = []

    for sheet in wb.worksheets:
        rows_iter = sheet.iter_rows(values_only=True)
        try:
            headers = [str(h) if h is not None else "" for h in next(rows_iter)]
        except StopIteration:
            continue

        sheet_label = f'{filename} — sheet "{sheet.title}"'
        for i, row in enumerate(rows_iter, start=2):
            if i - 1 > MAX_ROWS:
                break
            chunk = _row_to_chunk(sheet_label, i, headers, list(row))
            if chunk:
                chunks.append(chunk)

    return chunks


def parse_tabular_to_chunks(filename: str, data: bytes) -> list[str]:
    ext = extension_of(filename)
    if ext == ".csv":
        return _parse_csv_bytes(filename, data, delimiter=",")
    if ext == ".tsv":
        return _parse_csv_bytes(filename, data, delimiter="\t")
    if ext == ".xlsx":
        return _parse_xlsx_bytes(filename, data)
    raise ValueError(f"Not a tabular file type: {filename!r}")
